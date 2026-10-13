#!/usr/bin/env python3
"""Fill one TAP PDF per nonempty personnel row in an XLSX workbook."""

from __future__ import annotations

import argparse
import html
import re
import sys
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pypdf import PdfReader, PdfWriter


RANK_GRADES = {
    "PVT": "E1", "PV1": "E1", "PRIVATE": "E1",
    "PV2": "E2", "PRIVATE2": "E2", "PRIVATESECONDCLASS": "E2",
    "PFC": "E3", "PRIVATEFIRSTCLASS": "E3",
    "SPC": "E4", "SPECIALIST": "E4", "CPL": "E4", "CORPORAL": "E4",
    "SGT": "E5", "SERGEANT": "E5",
    "SSG": "E6", "STAFFSERGEANT": "E6",
    "SFC": "E7", "SERGEANTFIRSTCLASS": "E7",
    "MSG": "E8", "MASTERSERGEANT": "E8",
    "1SG": "E8", "FIRSTSERGEANT": "E8",
    "SGM": "E9", "SERGEANTMAJOR": "E9",
    "CSM": "E9", "COMMANDSERGEANTMAJOR": "E9",
    "SMA": "E9", "SERGEANTMAJOROFTHEARMY": "E9",
}


def normalize(value: object) -> str:
    """Ignore punctuation, HTML spacing, and case in column names/ranks."""
    text = re.sub(r"<[^>]*>", " ", html.unescape(str(value or "")))
    return re.sub(r"[^a-z0-9]", "", text.lower())


def grade_for_rank(rank: str) -> str:
    key = normalize(rank).upper()
    if re.fullmatch(r"E[1-9]", key):
        return key
    return RANK_GRADES.get(key, "")


def cell_text(cell: object) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, int) or value.is_integer():
            text = str(int(value))
            # Honor Excel's simple zero padding (e.g. a ten-digit DoD ID).
            if re.fullmatch(r"0+", cell.number_format or "") and value >= 0:
                text = text.zfill(len(cell.number_format))
            return text
    return str(value).strip()


def date_text(cell: object, epoch: datetime) -> str:
    if cell is None or cell.value is None:
        return ""
    value = cell.value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            converted = from_excel(value, epoch)
            if isinstance(converted, (date, datetime)):
                return converted.strftime("%m/%d/%Y")
        except (ValueError, OverflowError):
            pass
    text = cell_text(cell)
    for fmt in ("%m/%d/%Y", "%m-%d-%Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).strftime("%m/%d/%Y")
        except ValueError:
            continue
    return text  # Preserve unrecognized data rather than stopping the batch.


def split_name(name: str) -> tuple[str, str, str]:
    if "," not in name:
        # The expected format is Last, First M; don't guess a different order.
        return name.strip(), "", ""
    last, given = name.split(",", 1)
    parts = given.split()
    middle = ""
    if len(parts) > 1 and re.fullmatch(r"[^\W\d_]\.?", parts[-1], re.UNICODE):
        middle = parts.pop().rstrip(".")
    return last.strip(), " ".join(parts), middle


def fields_for_row(row: dict, epoch: datetime) -> dict[str, str]:
    def get_cell(*names: str):
        return next((row[normalize(n)] for n in names
                     if normalize(n) in row and cell_text(row[normalize(n)])), None)

    def text(*names: str) -> str:
        cell = get_cell(*names)
        return cell_text(cell) if cell is not None else ""

    last, first, middle = split_name(text("Name: Last, First M", "Name"))
    rank = text("Rank")
    separation = date_text(get_cell("Separation Date", "REFRAD"), epoch)
    return {
        "Last_Name": last,
        "First_Name": first,
        "Middle_Initial": middle,
        "DoD_ID": text("DOD ID", "DoD ID#"),
        "Date_of_Birth": date_text(get_cell("DOB", "Date of Birth"), epoch),
        "Cell_Phone": text("Contact Number", "Cell Phone"),
        "Gender": text("Gender"),
        "Rank": rank,
        "Grade": grade_for_rank(rank),
        "Component": "RC",
        "Unit_Code_UIC": "WS64AA",
        "Service_Affiliation": "Army",
        "Work_Email": text("Personal Email"),
        "Post_Transition_Location": text("Home of Record"),
        # These optional columns can be added to the supplied roster layout.
        "Separation_Date": separation,
        "Departure_Date": date_text(get_cell("Departure Date"), epoch) or separation,
        "MOS": text("MOS"),
        "Unit_Name": text("Unit Name"),
        "Salary": text("Salary", "Wanted Salary or current salary"),
        "Field_of_Interest": text("Field of Interest", "Job Title"),
    }


def output_path(directory: Path, last_name: str, row_number: int) -> Path:
    # Retain surnames/spaces, while preventing path traversal and invalid names.
    safe = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", last_name).strip(" .")
    safe = safe[:120].rstrip(" .") or f"unnamed_row_{row_number}"
    candidate = directory / f"{safe}_tapForm.pdf"
    counter = 2
    while candidate.exists():
        candidate = directory / f"{safe}_{counter}_tapForm.pdf"
        counter += 1
    return candidate


def fill_pdf(template: bytes, values: dict[str, str], destination: Path) -> None:
    writer = PdfWriter(clone_from=PdfReader(BytesIO(template)))
    writer.update_page_form_field_values(writer.pages[0], values, auto_regenerate=False)
    # Exclusive creation protects existing PDFs; remove partial output on failure.
    with destination.open("xb") as output:
        try:
            writer.write(output)
        except Exception:
            output.close()
            destination.unlink()
            raise


def fill_workbook(xlsx: Path, template: Path, directory: Path,
                  sheet_name: str | None = None) -> tuple[int, int]:
    template_bytes = template.read_bytes()
    fields = PdfReader(BytesIO(template_bytes)).get_fields() or {}
    if "Last_Name" not in fields:
        raise ValueError("The PDF template must contain the TAP Last_Name form field.")
    workbook = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if sheet_name is not None and sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} does not exist.")
        sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        directory.mkdir(parents=True, exist_ok=True)
        rows = iter(enumerate(sheet.iter_rows(), start=1))
        headers = None
        for _, cells in rows:
            if any(cell.value is not None for cell in cells):
                headers = [normalize(cell.value) for cell in cells]
                break
        if headers is None:
            raise ValueError("The worksheet is empty.")
        if len(set(h for h in headers if h)) != len([h for h in headers if h]):
            raise ValueError("The worksheet has duplicate column headers.")
        saved = failed = 0
        for row_number, cells in rows:
            row = {h: cell for h, cell in zip(headers, cells) if h}
            if not any(cell_text(cell) for cell in row.values()):
                continue
            values = fields_for_row(row, workbook.epoch)
            if values["Rank"] and not values["Grade"]:
                print(f"Row {row_number}: unknown/non-enlisted rank; grade left blank.",
                      file=sys.stderr)
            if not values["Last_Name"]:
                print(f"Row {row_number}: missing name; using a row-based filename.",
                      file=sys.stderr)
            destination = output_path(directory, values["Last_Name"], row_number)
            try:
                fill_pdf(template_bytes, values, destination)
            except Exception as error:
                failed += 1
                print(f"Row {row_number}: could not save PDF: {error}", file=sys.stderr)
                continue
            saved += 1
            print(f"Saved {destination}")
        if saved == 0 and failed == 0:
            raise ValueError("The worksheet contains no personnel rows.")
        return saved, failed
    finally:
        workbook.close()


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx", type=Path, help="Personnel XLSX; first nonempty row is the header")
    parser.add_argument("--template", type=Path,
                        default=Path(__file__).resolve().with_name("TAP_fillable.pdf"))
    parser.add_argument("--output-dir", type=Path, default=Path("filled_pdfs"))
    parser.add_argument("--sheet", help="Worksheet name (default: first worksheet)")
    args = parser.parse_args(argv)
    try:
        saved, failed = fill_workbook(args.xlsx, args.template, args.output_dir, args.sheet)
    except Exception as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1
    print(f"Completed: {saved} PDF(s) saved, {failed} failed.")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
